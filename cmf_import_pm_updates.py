#!/usr/bin/env python3
"""
cmf_import_pm_updates.py
========================
One-time import: pull PM updates from a downloaded file into the working file.

What it does
------------
1. Copies MAIN from the source file → working file (preserving images).
2. Reads green-highlighted rows from the source CALENDAR.
   Works with the OLD 7-sub-column layout (no hidden ID column) by matching
   WO# + due-date + station column back to a MAIN row.
   Also handles the NEW 8-sub-column layout (with hidden ID column).
3. Advances CURRENT_STEP in MAIN for every green-marked item.
4. Rebuilds CALENDAR, TODAY, and LOG in the working file.

Usage
-----
    python3 cmf_import_pm_updates.py  "<source_file.xlsx>"
"""

import os, sys, io, re
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image as XLImage
from datetime import datetime
from collections import defaultdict

_DIR      = os.path.dirname(os.path.abspath(__file__))
WORK_FILE = os.path.join(_DIR, "CMF WIP - Schedule.xlsx")

# ── Import schedule builder constants & functions ─────────────────────────────
sys.path.insert(0, _DIR)
import cmf_schedule_builder as sb

# ── Old-calendar layout constants ─────────────────────────────────────────────
OLD_SUB_N = 7   # STATUS|DUE|SHIP|COMPANY|WO#|QTY|PHOTO
OLD_BLOCK = 8   # OLD_SUB_N + 1 divider
# Offsets within each old block
OLD_OFF_STATUS  = 0
OLD_OFF_DUE     = 1
OLD_OFF_COMPANY = 3
OLD_OFF_WO      = 4

# Maps display name (upper) → station key
_NAME_TO_KEY = {v.upper(): k for k, v in sb.STATION_NAME.items()}


# ── Green detection (handles theme=6 AND explicit RGB) ───────────────────────
def _is_green(cell):
    try:
        fill = cell.fill
        if not fill or fill.fill_type != "solid":
            return False
        fg = fill.fgColor
        if fg.type == "theme":
            return int(fg.theme) == 6
        rgb = str(fg.rgb or "")
        if not rgb or rgb in ("00000000", "FF000000", "FFFFFFFF"):
            return False
        if len(rgb) == 8:
            r, g, b = int(rgb[2:4], 16), int(rgb[4:6], 16), int(rgb[6:8], 16)
        elif len(rgb) == 6:
            r, g, b = int(rgb[0:2], 16), int(rgb[2:4], 16), int(rgb[4:6], 16)
        else:
            return False
        return (g - b) > 30 and g > r and (r + g + b) // 3 < 200
    except Exception:
        return False


# ── Parse source CALENDAR for green completions ───────────────────────────────
def parse_green_from_calendar(ws_cal, ws_main_src):
    """
    Detect all green-highlighted rows in ws_cal.
    Supports both old layout (no ID col, CAL_BLOCK=8) and new layout (ID col, CAL_BLOCK=9).
    Returns list of {"main_row": int, "s_key": str}.
    """
    completions = []
    seen = set()

    # Detect layout by checking row 4 col 1 header
    hdr_c1 = str(ws_cal.cell(4, 1).value or "").strip()
    if hdr_c1 == "STATUS":
        # ── OLD layout: no hidden ID ──────────────────────────────────────
        print("  Detected OLD calendar layout (7 sub-cols)")
        block_size = OLD_BLOCK

        # Build station block map from row 3 header text
        block_to_skey = {}
        col = 1
        while col <= ws_cal.max_column:
            raw = str(ws_cal.cell(3, col).value or "").strip()
            clean = re.sub(r"\s*\(\d+\).*$", "", raw).strip().upper()
            s_key = _NAME_TO_KEY.get(clean)
            if s_key:
                block_to_skey[col] = s_key
            col += block_size

        print(f"  Station blocks: {len(block_to_skey)}")

        # Build MAIN lookup: (wo_norm, proc_col, due_str) → [row]
        lookup = defaultdict(list)
        for row in range(sb.DATA_START, ws_main_src.max_row + 1):
            wo = ws_main_src.cell(row, sb.COL_WO).value
            if not wo:
                continue
            wo_k = _norm_wo(wo)
            for c in range(12, 36):
                v = ws_main_src.cell(row, c).value
                if isinstance(v, datetime):
                    lookup[(wo_k, c, f"{v.month}/{v.day}")].append(row)

        # Scan data rows
        for row in range(5, ws_cal.max_row + 1):
            for sc, s_key in block_to_skey.items():
                block_green = any(_is_green(ws_cal.cell(row, sc + off))
                                  for off in range(OLD_SUB_N))
                if not block_green:
                    continue

                wo_raw  = ws_cal.cell(row, sc + OLD_OFF_WO).value
                due_str = str(ws_cal.cell(row, sc + OLD_OFF_DUE).value or "").strip()
                if not wo_raw or not due_str:
                    continue

                wo_k    = _norm_wo(wo_raw)
                s_def   = next((s for s in sb.STATIONS if s[0] == s_key), None)
                if not s_def:
                    continue

                for proc_col in s_def[4]:
                    for main_row in lookup.get((wo_k, proc_col, due_str), []):
                        key = (main_row, s_key)
                        if key not in seen:
                            seen.add(key)
                            completions.append({"main_row": main_row, "s_key": s_key})

    else:
        # ── NEW layout: hidden ID column at offset 0 ──────────────────────
        print("  Detected NEW calendar layout (8 sub-cols with hidden ID)")
        for row in range(5, ws_cal.max_row + 1):
            col = 1
            while col <= ws_cal.max_column:
                id_val = str(ws_cal.cell(row, col).value or "")
                if "|" in id_val:
                    block_green = any(_is_green(ws_cal.cell(row, col + off))
                                      for off in range(1, sb.CAL_SUB_N))
                    if block_green:
                        parts = id_val.split("|", 1)
                        if len(parts) == 2:
                            try:
                                key = (int(parts[0]), parts[1])
                                if key not in seen:
                                    seen.add(key)
                                    completions.append({"main_row": key[0], "s_key": key[1]})
                            except ValueError:
                                pass
                col += sb.CAL_BLOCK

    print(f"  Green completions found: {len(completions)}")
    for c in completions:
        wo  = ws_main_src.cell(c["main_row"], sb.COL_WO).value
        cmp = ws_main_src.cell(c["main_row"], sb.COL_COMPANY).value or ""
        print(f"    WO {wo}  {cmp}  [{sb.STATION_NAME.get(c['s_key'], c['s_key'])}]")

    return completions


def _norm_wo(val):
    s = str(val).strip()
    try:
        f = float(s)
        if f == int(f):
            s = str(int(f))
    except ValueError:
        pass
    return s


# ── Copy MAIN from source → destination workbook ──────────────────────────────
def copy_main_sheet(wb_src, wb_dst):
    """Replace MAIN in wb_dst with MAIN from wb_src, including images."""
    ws_src = wb_src["MAIN"]

    if "MAIN" in wb_dst.sheetnames:
        del wb_dst["MAIN"]
    ws_dst = wb_dst.create_sheet("MAIN", 0)

    # Copy cell values, styles, dimensions
    for row in ws_src.iter_rows():
        for cell in row:
            new_cell = ws_dst.cell(row=cell.row, column=cell.column, value=cell.value)
            if cell.has_style:
                new_cell.font      = cell.font.copy()
                new_cell.fill      = cell.fill.copy()
                new_cell.border    = cell.border.copy()
                new_cell.alignment = cell.alignment.copy()
                new_cell.number_format = cell.number_format

    # Column widths, row heights
    for col_letter, cd in ws_src.column_dimensions.items():
        ws_dst.column_dimensions[col_letter].width = cd.width
    for row_idx, rd in ws_src.row_dimensions.items():
        ws_dst.row_dimensions[row_idx].height = rd.height

    # Merged cells
    for mc in ws_src.merged_cells.ranges:
        ws_dst.merge_cells(str(mc))

    # Freeze panes — always lock just the two header rows regardless of source view
    ws_dst.freeze_panes = "A3"
    if ws_src.auto_filter.ref:
        ws_dst.auto_filter.ref = ws_src.auto_filter.ref

    # Data validations
    for dv in ws_src.data_validations.dataValidation:
        ws_dst.add_data_validation(dv)

    # Conditional formatting
    for rng, rules in ws_src.conditional_formatting._cf_rules.items():
        for rule in rules:
            ws_dst.conditional_formatting.add(rng, rule)

    # Images
    img_count = 0
    for img in ws_src._images:
        try:
            ref = img.ref
            if hasattr(ref, "read"):
                ref.seek(0)
                raw = ref.read()
                ref.seek(0)
            else:
                raw = img._data()
            if not raw:
                continue
            anchor_from = img.anchor._from
            sb.copy_image(raw, ws_dst, anchor_from.col + 1, anchor_from.row + 1)
            img_count += 1
        except Exception:
            pass

    print(f"  MAIN copied: {ws_src.max_row} rows, {ws_src.max_column} cols, {img_count} images")
    return ws_dst


# ── Main ──────────────────────────────────────────────────────────────────────
if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Usage: python3 cmf_import_pm_updates.py \"<source_file.xlsx>\"")
        raise SystemExit(1)

    src_path = sys.argv[1]
    if not os.path.exists(src_path):
        print(f"ERROR: Source file not found: {src_path}")
        raise SystemExit(1)

    print(f"Source:  {src_path}")
    print(f"Working: {WORK_FILE}")
    print()

    # Load both files
    print("Loading source file ...")
    wb_src = openpyxl.load_workbook(src_path)
    sb.protect_workbook_images(wb_src)

    print("Loading working file ...")
    wb_work = openpyxl.load_workbook(WORK_FILE) if os.path.exists(WORK_FILE) else openpyxl.Workbook()
    sb.protect_workbook_images(wb_work)

    # Step 1: Parse green completions BEFORE overwriting MAIN
    print()
    print("Scanning source CALENDAR for green completions ...")
    ws_main_src = wb_src["MAIN"] if "MAIN" in wb_src.sheetnames else wb_src.active
    ws_cal_src  = wb_src["CALENDAR"] if "CALENDAR" in wb_src.sheetnames else None

    completions = []
    if ws_cal_src:
        completions = parse_green_from_calendar(ws_cal_src, ws_main_src)
    else:
        print("  No CALENDAR sheet in source — skipping green sync")

    # Step 2: Copy MAIN from source → working file
    print()
    print("Copying MAIN from source → working file ...")
    sb.ensure_ship_vendor_col(ws_main_src)
    ws_main_new = copy_main_sheet(wb_src, wb_work)

    # Step 3: Apply completions to new MAIN
    if completions:
        print()
        print("Advancing CURRENT_STEP in MAIN ...")
        records = sb.apply_completions(ws_main_new, completions)
        print()
        print("Recording to LOG ...")
        sb.record_completions_to_log(wb_work, records)
    else:
        print()
        print("  No green completions — MAIN current steps unchanged")

    # Step 4: Re-parse MAIN and rebuild CALENDAR + TODAY
    print()
    print("Parsing updated MAIN ...")
    entries, row_to_img = sb.parse_main(ws_main_new)
    print(f"  {len(entries)} process-due entries  |  {len(row_to_img)} screenshots")

    print("Building CALENDAR ...")
    sb.build_calendar(wb_work, entries, row_to_img)

    print("Building TODAY (Kanban) ...")
    sb.build_today(wb_work, entries, row_to_img)

    # Step 5: Sheet ordering and save
    for name in ["LOG", "TODAY", "CALENDAR", "MAIN"]:
        if name in wb_work.sheetnames:
            wb_work.move_sheet(name, offset=-len(wb_work.sheetnames))

    print()
    print(f"Saving → {WORK_FILE}")
    wb_work.save(WORK_FILE)
    print("Done ✓")
