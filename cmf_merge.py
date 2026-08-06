#!/usr/bin/env python3
"""
cmf_merge.py  —  Run ONCE to merge boss's updates with engineer's routing
Inputs:
  CMF WIP (1).xlsx           ← boss's latest orders/parts (old format)
  CMF WIP - Schedule (1).xlsx ← engineer's routing/process dates (new format)
Output:
  CMF WIP - Schedule.xlsx    ← merged result (overwrites current working file)

Matching logic:
  Primary:   WO# + Part#  (exact match)
  Secondary: WO# only     (for WO-level routing with no specific part#)
"""

import os, io, re, sys, shutil
from collections import OrderedDict
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.units import pixels_to_EMU
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.drawing.image import Image as XLImage
from datetime import datetime

_DIR      = os.path.dirname(os.path.abspath(__file__))
OUT_FILE  = os.path.join(_DIR, 'CMF WIP - Schedule.xlsx')
REVIEW_FILE = os.path.join(_DIR, 'Merge Review.xlsx')
DRY_RUN   = "--dry-run" in sys.argv
WIP_RE    = re.compile(r"^CMF WIP(?: \((\d+)\))?\.xlsx$")
SCH_RE    = re.compile(r"^CMF WIP - Schedule(?: \((\d+)\))?\.xlsx$")

# ── Palette ───────────────────────────────────────────────────────────────────
C_NAVY="1F3864"; C_GOLD="FFD966"; C_STEEL_DK="1F4E79"; C_BLUE_LT="DEEAF1"
C_WHITE="FFFFFF"; C_LGRAY="F2F2F2"; C_DGRAY="404040"
C_YELLOW="FFF2CC"; C_RED_LT="FFD7D7"

def _fill(h): return PatternFill(start_color=h,end_color=h,fill_type="solid")
def _font(bold=False,color=C_DGRAY,size=10): return Font(name="Calibri",bold=bold,color=color,size=size)
def _border(c="CCCCCC"): s=Side(style="thin",color=c); return Border(left=s,right=s,top=s,bottom=s)
def _align(h="center",v="center",wrap=False): return Alignment(horizontal=h,vertical=v,wrap_text=wrap)

IMG_W=85; IMG_H=48; COL_I_W=13; PART_ROW_H=52
IMG_PAD_X=2; IMG_PAD_Y=2
OUT_SHIP_VENDOR_COL = 35
OUT_DELIVERY_COL    = 36
OUT_CURRENT_COL     = 37
OUT_STATUS_COL      = 38
OUT_NOTES_COL       = 39

COLS = [
    ("A","WO #",8,False),("B","PO #",17,False),("C","COMPANY",20,False),
    ("D","JOB NAME",28,False),("E","CUST SHIP",12,False),
    ("F","PART #",16,False),("G","DESCRIPTION",26,False),
    ("H","QTY",7,False),("I","SCREENSHOT",COL_I_W,False),
    ("J","MATERIAL",10,False),("K","THICKNESS",9,False),
    ("L","MATL",6,True),("M","ENG",6,True),
    ("N","LASER",6,True),("O","LASER\n(O)",6,True),
    ("P","TUBE\nLSR(O)",6,True),("Q","SAW",6,True),
    ("R","BAND\n(O)",6,True),("S","BEND",6,True),
    ("T","CLEAN",6,True),("U","CSK",6,True),
    ("V","DRILL",6,True),("W","TAP",6,True),
    ("X","GRIND",6,True),("Y","WELD",6,True),
    ("Z","MACH\n(O)",6,True),("AA","PLAT\n(O)",6,True),
    ("AB","PAINT",6,True),("AC","PAINT\n(O)",6,True),
    ("AD","SPEC",6,True),("AE","SPEC\n(O)",6,True),
    ("AF","W.JOB\n(O)",6,True),("AG","HW",6,True),
    ("AH","ASSM",6,True),("AI","SHIP TO\nVENDOR",8,True),
    ("AJ","DELIVERY\nDATE",8,True),("AK","CURRENT\nSTEP",18,False),
    ("AL","STATUS",12,False),("AM","NOTES",35,False),
]
NCOLS=len(COLS); LAST="AM"; HDR=2; DATA=3
BLUE_FIRST=12; BLUE_LAST=OUT_DELIVERY_COL

STEP_LIST=(
    "MATERIALS,ENGINEERING,LASER CUT,LASER CUT (O),TUBE LASER (O),"
    "SAW,BANDSAW (O),BEND,CLEAN,CSK,DRILL,TAPPING,GRIND,WELD,"
    "MACHINE (O),PLATING (O),PAINT,PAINT (O),"
    "SPECIAL,SPECIAL (O),WHOLE JOB (O),HARDWARE,ASSEMBLY,SHIP TO VENDOR,SHIP,RECEIVING,"
    "COMPLETE,ON HOLD"
)


# ── Image helpers ─────────────────────────────────────────────────────────────
class _UnclosableBytesIO(io.BytesIO):
    def close(self): pass

def _set_two_cell_anchor(img, col_1idx, row_1idx, width_px=IMG_W, height_px=IMG_H, pad_x=IMG_PAD_X, pad_y=IMG_PAD_Y):
    from openpyxl.drawing.spreadsheet_drawing import TwoCellAnchor, AnchorMarker
    c0, r0 = col_1idx - 1, row_1idx - 1
    anchor = TwoCellAnchor(editAs='twoCell')
    anchor._from = AnchorMarker(col=c0, colOff=pixels_to_EMU(pad_x), row=r0, rowOff=pixels_to_EMU(pad_y))
    anchor.to = AnchorMarker(col=c0, colOff=pixels_to_EMU(pad_x + width_px), row=r0, rowOff=pixels_to_EMU(pad_y + height_px))
    img.anchor = anchor

def _copy_image(src_img, dest_ws, col_1idx, row_1idx):
    try:
        raw=src_img._data(); buf=_UnclosableBytesIO(raw); img=XLImage(buf)
        _set_two_cell_anchor(img, col_1idx, row_1idx)
        dest_ws.add_image(img); return True
    except Exception: return False


def _norm_key(val):
    """Normalise a WO# or part# to a consistent string for matching."""
    if val is None: return None
    s = str(val).strip()
    # If it looks like a float (e.g. "4418.0"), convert to int string
    try:
        f = float(s)
        if f == int(f): s = str(int(f))
    except ValueError: pass
    return s or None


def _backup_file(path):
    """Copy path to a timestamped .bak before it gets overwritten.

    Every merge/refresh rewrites the working file in place. Without this the
    engineer's routing is unrecoverable if a run picks the wrong input or
    misparses a row.
    """
    if not os.path.exists(path):
        return None
    backup_dir = os.path.join(_DIR, "backups")
    os.makedirs(backup_dir, exist_ok=True)
    stamp = datetime.now().strftime("%Y-%m-%d_%H%M%S")
    base, ext = os.path.splitext(os.path.basename(path))
    dest = os.path.join(backup_dir, f"{base}__{stamp}{ext}")
    shutil.copy2(path, dest)
    print(f"  Backup: backups/{os.path.basename(dest)}")
    return dest


def _pick_latest(pattern, label):
    """Pick the most recently modified file matching pattern.

    Selection is by modification time, NOT by the "(N)" suffix. The old
    behaviour preferred the highest (N), so a fresh re-download named
    "CMF WIP.xlsx" lost to a stale "CMF WIP (1).xlsx" from a previous week
    and the merge silently ran on old orders.
    """
    candidates = []
    for name in os.listdir(_DIR):
        if not pattern.match(name):
            continue
        path = os.path.join(_DIR, name)
        candidates.append((os.path.getmtime(path), path))

    if not candidates:
        raise FileNotFoundError(f"No files matched {pattern.pattern}")

    candidates.sort(reverse=True)
    chosen_mtime, chosen = candidates[0]

    print(f"  {label} candidates:")
    for mtime, path in candidates:
        mark = "USING >" if path == chosen else "       "
        age = (datetime.now() - datetime.fromtimestamp(mtime)).total_seconds() / 3600
        print(f"    {mark} {os.path.basename(path):<38} modified {datetime.fromtimestamp(mtime):%Y-%m-%d %H:%M} ({age:.1f}h ago)")

    age_hours = (datetime.now() - datetime.fromtimestamp(chosen_mtime)).total_seconds() / 3600
    if age_hours > 48:
        print(f"    !! WARNING: newest {label} file is {age_hours/24:.1f} days old — is this the export you meant?")
    if len(candidates) > 1 and abs(candidates[0][0] - candidates[1][0]) < 120:
        print(f"    !! WARNING: two {label} files modified within 2 minutes — delete the one you don't want and re-run.")

    return chosen


def resolve_input_files():
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    if len(args) >= 2:
        return os.path.abspath(args[0]), os.path.abspath(args[1])
    return _pick_latest(WIP_RE, "WIP"), _pick_latest(SCH_RE, "Schedule")


def detect_schedule_layout(ws):
    hdr_35 = str(ws.cell(HDR, 35).value or "").replace("\n", " ").strip().upper()
    has_ship_vendor = "SHIP TO VENDOR" in hdr_35
    return {
        "delivery_col": 36 if has_ship_vendor else 35,
        "current_col": 37 if has_ship_vendor else 36,
        "status_col": 38 if has_ship_vendor else 37,
        "notes_col": 39 if has_ship_vendor else 38,
    }


# ── Build routing lookup from Schedule(1) ────────────────────────────────────
def build_routing_lookup(ws_sch, layout):
    """
    Returns:
      by_part: {(wo_str, part_str): {col: date}}   ← exact WO+Part match
      by_wo:   {wo_str: {col: date}}                ← WO-level (no part#)
    """
    by_part = {}
    by_wo   = {}

    for row in range(DATA, ws_sch.max_row + 1):
        wo_raw  = ws_sch.cell(row, 1).value
        pno_raw = ws_sch.cell(row, 6).value
        if not wo_raw: continue

        routing = {}
        for col in range(12, layout["delivery_col"] + 1):
            v = ws_sch.cell(row, col).value
            if v and isinstance(v, datetime):
                out_col = OUT_DELIVERY_COL if col == layout["delivery_col"] else col
                routing[out_col] = v

        if not routing: continue

        wo_k  = _norm_key(wo_raw)
        pno_k = _norm_key(pno_raw)

        if pno_k:
            # Merge (later rows win for same key)
            key = (wo_k, pno_k)
            existing = by_part.get(key, {})
            existing.update(routing)
            by_part[key] = existing
        else:
            # WO-level routing
            existing = by_wo.get(wo_k, {})
            existing.update(routing)
            by_wo[wo_k] = existing

    print(f"  Routing lookup: {len(by_part)} part-level + {len(by_wo)} WO-level entries")
    return by_part, by_wo


def get_routing(wo, pno, by_part, by_wo):
    """Return merged routing dict for this WO+Part (part-level wins over WO-level)."""
    wo_k  = _norm_key(wo)
    pno_k = _norm_key(pno)
    result = dict(by_wo.get(wo_k, {}))           # start with WO-level
    if pno_k:
        result.update(by_part.get((wo_k, pno_k), {}))  # part-level overrides
    return result


def _merge_schedule_state(existing, incoming):
    if not existing:
        existing = {"routing": {}, "current": None, "status": None, "notes": None}
    existing["routing"].update(incoming.get("routing", {}))
    for key in ("current", "status", "notes"):
        val = incoming.get(key)
        if val not in (None, ""):
            existing[key] = val
    return existing


def build_schedule_state_lookup(ws_sch, layout):
    by_part = {}
    by_wo = {}
    ship_by_wo = {}
    summary_by_wo = {}
    current_wo = None

    for row in range(DATA, ws_sch.max_row + 1):
        wo_raw = ws_sch.cell(row, 1).value
        wo_k = _norm_key(wo_raw)
        if wo_k:
            current_wo = wo_k
        if not current_wo:
            continue

        po = ws_sch.cell(row, 2).value
        company = ws_sch.cell(row, 3).value
        job_name = ws_sch.cell(row, 4).value
        cust_ship = ws_sch.cell(row, 5).value
        part = ws_sch.cell(row, 6).value
        desc = ws_sch.cell(row, 7).value
        qty = ws_sch.cell(row, 8).value

        summary = summary_by_wo.setdefault(current_wo, {
            "wo": current_wo,
            "po": None,
            "company": None,
            "job_name": None,
            "cust_ship": None,
        })
        if po not in (None, ""):
            summary["po"] = po
        if company not in (None, ""):
            summary["company"] = company
        if job_name not in (None, ""):
            summary["job_name"] = job_name
        if cust_ship not in (None, ""):
            summary["cust_ship"] = cust_ship

        if part is None and desc is None and qty is None:
            if isinstance(cust_ship, datetime):
                ship_by_wo[current_wo] = cust_ship
            continue

        state = {
            "routing": {},
            "current": ws_sch.cell(row, layout["current_col"]).value,
            "status": ws_sch.cell(row, layout["status_col"]).value,
            "notes": ws_sch.cell(row, layout["notes_col"]).value,
        }
        for col in range(12, layout["delivery_col"] + 1):
            v = ws_sch.cell(row, col).value
            if isinstance(v, datetime):
                out_col = OUT_DELIVERY_COL if col == layout["delivery_col"] else col
                state["routing"][out_col] = v

        pno_k = _norm_key(part)
        if pno_k:
            key = (current_wo, pno_k)
            by_part[key] = _merge_schedule_state(by_part.get(key), state)
        else:
            by_wo[current_wo] = _merge_schedule_state(by_wo.get(current_wo), state)

    print(f"  Schedule state: {len(summary_by_wo)} WOs | {len(by_part)} part-level + {len(by_wo)} WO-level entries")
    return by_part, by_wo, ship_by_wo, summary_by_wo


def get_schedule_state(wo, pno, by_part, by_wo):
    wo_k = _norm_key(wo)
    pno_k = _norm_key(pno)
    result = {"routing": {}, "current": None, "status": None, "notes": None}

    wo_state = by_wo.get(wo_k)
    if wo_state:
        result = _merge_schedule_state(result, wo_state)

    if pno_k:
        part_state = by_part.get((wo_k, pno_k))
        if part_state:
            result = _merge_schedule_state(result, part_state)

    return result


# ── Sheet setup (same as migrate script) ─────────────────────────────────────
def _setup_sheet(ws):
    ws.row_dimensions[1].height = 30
    ws.merge_cells(f"A1:{LAST}1")
    c=ws["A1"]
    c.value=f"CMF  WORK IN PROGRESS   —   Updated {datetime.now().strftime('%B %d, %Y')}"
    c.font=_font(bold=True,color=C_WHITE,size=14)
    c.fill=_fill(C_NAVY); c.alignment=_align()

    ws.row_dimensions[HDR].height = 42
    for i,(letter,label,width,is_blue) in enumerate(COLS,1):
        ws.column_dimensions[letter].width = width
        c=ws.cell(HDR,i,label); c.border=_border(); c.alignment=_align(wrap=True)
        if is_blue: c.font=_font(bold=True,size=8,color=C_STEEL_DK); c.fill=_fill(C_BLUE_LT)
        else:       c.font=_font(bold=True,size=8,color=C_WHITE);     c.fill=_fill(C_NAVY)


def _write_header(ws, row, wo, po, company, job_name, cust_ship):
    ws.row_dimensions[row].height = 24
    vals=[wo,po,company,job_name,cust_ship]+[None]*(NCOLS-5)
    for i,val in enumerate(vals,1):
        c=ws.cell(row,i,val); c.fill=_fill(C_NAVY); c.border=_border("2E4D7B")
        c.alignment=_align(h="left")
        c.font=_font(bold=True,color=(C_GOLD if i==1 else C_WHITE),size=11)
        if isinstance(val,datetime): c.number_format="M/D/YY"


def _write_part(ws, row, wo, po, company, pno, desc, qty, mat, thk,
                current, status, notes, routing, alt):
    ws.row_dimensions[row].height = PART_ROW_H
    bg = C_LGRAY if alt else C_WHITE

    # Build vals: 39 cols total
    # A-E(5), F-K(6), L-S(8), T-Y(6), Z-AJ(11), AK-AM(3)
    vals = [
        wo, po, company, None, None,                    # A-E
        pno, desc, qty, None, mat, thk,                 # F-K  (I=screenshot)
        None,None,None,None,None,None,None,None,        # L-S  (8)
        None,None,None,None,None,None,                  # T-Y  (6)
        None,None,None,None,None,None,None,None,None,None,None,  # Z-AJ (11)
        current, status, notes,                         # AK-AM
    ]
    assert len(vals)==NCOLS, f"vals len {len(vals)} != {NCOLS}"

    # Apply routing dates into the correct positions
    for col_idx, date_val in routing.items():
        list_pos = col_idx - 1   # 1-indexed col → 0-indexed list position
        if 0 <= list_pos < NCOLS:
            vals[list_pos] = date_val

    for i,val in enumerate(vals,1):
        _,_,_,is_blue=COLS[i-1]
        c=ws.cell(row,i,val); c.border=_border()
        c.fill=_fill(C_BLUE_LT if (is_blue and val is None) else
                     (C_WHITE  if is_blue else bg))
        c.font=_font(bold=(i in (1,2,3)),size=10)
        c.alignment=_align(h="left" if i in (3,6,7,38) else "center")
        if isinstance(val,datetime): c.number_format="M/D/YY"


def _apply_cf(ws, last_row):
    fc=get_column_letter(BLUE_FIRST); lc=get_column_letter(BLUE_LAST)
    rng=f"{fc}{DATA}:{lc}{last_row}"
    ws.conditional_formatting.add(rng, FormulaRule(
        formula=[f"AND({fc}{DATA}<>\"\",ISNUMBER({fc}{DATA}),{fc}{DATA}<TODAY())"],
        fill=_fill(C_RED_LT),font=Font(color="CC0000",bold=True)))
    ws.conditional_formatting.add(rng, FormulaRule(
        formula=[f"AND({fc}{DATA}<>\"\",ISNUMBER({fc}{DATA}),{fc}{DATA}>=TODAY(),{fc}{DATA}<=TODAY()+3)"],
        fill=_fill(C_YELLOW),font=Font(color="806000",bold=True)))


def _apply_dropdown(ws, last_row):
    dv=DataValidation(type="list",formula1=f'"{STEP_LIST}"',allow_blank=True,showErrorMessage=False)
    dv.sqref=f"AK{DATA}:AK{last_row}"; ws.add_data_validation(dv)


def _display_date(val):
    if isinstance(val, datetime):
        return f"{val.month}/{val.day}/{val.year}"
    return "" if val in (None, "") else str(val)


def _write_change_section(ws, row, title, headers, rows):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(headers))
    cell = ws.cell(row, 1, title)
    cell.font = _font(bold=True, color=C_WHITE, size=11)
    cell.fill = _fill(C_STEEL_DK)
    cell.alignment = _align(h="left")
    for col in range(1, len(headers) + 1):
        ws.cell(row, col).border = _border("2E4D7B")
    row += 1

    for idx, header in enumerate(headers, 1):
        c = ws.cell(row, idx, header)
        c.font = _font(bold=True, size=9, color=C_STEEL_DK)
        c.fill = _fill(C_BLUE_LT)
        c.border = _border()
        c.alignment = _align()
    row += 1

    if not rows:
        ws.cell(row, 1, "None")
        ws.cell(row, 1).border = _border()
        row += 2
        return row

    alt = False
    for values in rows:
        alt = not alt
        bg = C_LGRAY if alt else C_WHITE
        for idx, value in enumerate(values, 1):
            c = ws.cell(row, idx, value)
            c.border = _border()
            c.fill = _fill(bg)
            c.alignment = _align(h="left" if idx > 1 else "center")
        row += 1
    return row + 1


def build_change_log_sheet(wb, wip_path, sch_path, wip_summary_by_wo, sch_summary_by_wo,
                           new_wos, removed_wos, ship_date_updates,
                           routing_counts_by_wo, missing_routing_by_wo):
    if "MERGE CHANGES" in wb.sheetnames:
        del wb["MERGE CHANGES"]
    ws = wb.create_sheet("MERGE CHANGES")
    widths = [12, 18, 28, 30, 18, 18]
    for idx, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = "A4"

    ws.merge_cells("A1:F1")
    c = ws["A1"]
    c.value = f"CMF  MERGE CHANGES   —   Updated {datetime.now().strftime('%B %d, %Y %I:%M %p')}"
    c.font = _font(bold=True, color=C_WHITE, size=14)
    c.fill = _fill(C_NAVY)
    c.alignment = _align()

    meta_rows = [
        ("Google/WIP source", os.path.basename(wip_path)),
        ("Excel/Schedule source", os.path.basename(sch_path)),
        ("Merged output", os.path.basename(OUT_FILE)),
    ]
    row = 3
    for label, value in meta_rows:
        ws.cell(row, 1, label).font = _font(bold=True)
        ws.cell(row, 1).fill = _fill(C_BLUE_LT)
        ws.cell(row, 1).border = _border()
        ws.cell(row, 2, value).border = _border()
        row += 1

    summary_rows = [
        ("WOs in Google WIP", len(wip_summary_by_wo)),
        ("WOs in Excel schedule", len(sch_summary_by_wo)),
        ("New WOs added", len(new_wos)),
        ("Removed WOs", len(removed_wos)),
        ("Ship date updates from Excel", len(ship_date_updates)),
        ("WOs with routing carried over", len(routing_counts_by_wo)),
        ("WOs needing routing work", len(missing_routing_by_wo)),
    ]
    row = _write_change_section(ws, row + 1, "Summary", ["Metric", "Value"], summary_rows)

    new_rows = []
    for wo in sorted(new_wos, key=lambda v: int(v) if str(v).isdigit() else str(v)):
        s = wip_summary_by_wo.get(wo, {})
        new_rows.append([wo, s.get("po", ""), s.get("company", ""), s.get("job_name", ""), _display_date(s.get("cust_ship")), s.get("part_count", 0)])
    row = _write_change_section(ws, row, "New WOs Added From Google", ["WO #", "PO #", "Company", "Job Name", "Ship Date", "Parts"], new_rows)

    removed_rows = []
    for wo in sorted(removed_wos, key=lambda v: int(v) if str(v).isdigit() else str(v)):
        s = sch_summary_by_wo.get(wo, {})
        removed_rows.append([wo, s.get("po", ""), s.get("company", ""), s.get("job_name", ""), _display_date(s.get("cust_ship")), "Removed"])
    row = _write_change_section(ws, row, "Removed WOs (Missing From Google)", ["WO #", "PO #", "Company", "Job Name", "Prior Ship", "Status"], removed_rows)

    ship_rows = []
    for change in sorted(ship_date_updates, key=lambda item: int(item["wo"]) if str(item["wo"]).isdigit() else str(item["wo"])):
        ship_rows.append([change["wo"], change["po"], change["company"], change["job_name"], _display_date(change["google_ship"]), _display_date(change["excel_ship"])])
    row = _write_change_section(ws, row, "Ship Date Updates Kept From Excel", ["WO #", "PO #", "Company", "Job Name", "Google Ship", "Excel Ship"], ship_rows)

    routing_rows = []
    for wo, count in sorted(routing_counts_by_wo.items(), key=lambda item: int(item[0]) if str(item[0]).isdigit() else str(item[0])):
        s = wip_summary_by_wo.get(wo) or sch_summary_by_wo.get(wo) or {}
        routing_rows.append([wo, s.get("company", ""), s.get("job_name", ""), count, "", "Routing carried over"])
    row = _write_change_section(ws, row, "Routing Carried Over From Excel", ["WO #", "Company", "Job Name", "Part Rows", "", "Status"], routing_rows)

    missing_rows = []
    for wo, count in sorted(missing_routing_by_wo.items(), key=lambda item: int(item[0]) if str(item[0]).isdigit() else str(item[0])):
        s = wip_summary_by_wo.get(wo, {})
        missing_rows.append([wo, s.get("company", ""), s.get("job_name", ""), count, _display_date(s.get("cust_ship")), "Needs routing"])
    _write_change_section(ws, row, "Part Rows Needing Routing", ["WO #", "Company", "Job Name", "Part Rows", "Ship Date", "Status"], missing_rows)


# ── Main merge ────────────────────────────────────────────────────────────────
def merge():
    wip_path, sch_path = resolve_input_files()
    print(f"Loading WIP:         {wip_path}")
    wb_wip = openpyxl.load_workbook(wip_path)
    ws_wip = wb_wip['MAIN']

    print(f"Loading Schedule:    {sch_path}")
    wb_sch = openpyxl.load_workbook(sch_path)
    ws_sch = wb_sch['MAIN']
    sch_layout = detect_schedule_layout(ws_sch)

    # Build image maps from BOTH sources
    # WIP(1) images are primary (more up-to-date), by old_excel_row
    wip_img_map = {}
    for img in ws_wip._images:
        try: wip_img_map[img.anchor._from.row + 1] = img
        except Exception: pass

    # Schedule(1) images keyed by (WO#, Part#) for secondary lookup
    sch_img_by_key = {}
    for row in range(DATA, ws_sch.max_row + 1):
        if not ws_sch._images: break
    for img in ws_sch._images:
        try:
            er = img.anchor._from.row + 1
            wo  = _norm_key(ws_sch.cell(er, 1).value)
            pno = _norm_key(ws_sch.cell(er, 6).value)
            if wo: sch_img_by_key[(wo, pno)] = img
        except Exception: pass

    print(f"  WIP images: {len(wip_img_map)}  |  Schedule images: {len(sch_img_by_key)}")

    # Build routing + schedule-state lookup from Schedule source
    print("Building schedule lookup ...")
    by_part, by_wo = build_routing_lookup(ws_sch, sch_layout)
    state_by_part, state_by_wo, ship_by_wo, sch_summary_by_wo = build_schedule_state_lookup(ws_sch, sch_layout)

    # Pre-group WIP(1) rows by WO (preserve order)
    wo_groups = OrderedDict()
    in_row_num = 2
    for row_data in ws_wip.iter_rows(min_row=2, values_only=True):
        wo = row_data[8]    # old col I = WO NO (0-indexed 8 = 1-indexed 9)
        if wo:
            if wo not in wo_groups: wo_groups[wo] = []
            wo_groups[wo].append((in_row_num, row_data))
        in_row_num += 1

    print(f"  {len(wo_groups)} WOs from Google/WIP")

    wip_summary_by_wo = {}
    for wo, rows in wo_groups.items():
        first_er, first_row = rows[0]
        first_is_summary = (first_row[0] is None and first_row[1] is None)
        job_name = (first_row[10] or first_row[9]) if first_is_summary else None
        if job_name:
            job_name = str(job_name).strip()
        company = first_row[6]
        po = first_row[7]
        due = first_row[3]
        part_rows = rows[1:] if first_is_summary else rows
        if not part_rows:
            part_rows = [(first_er, first_row)]
        wip_summary_by_wo[_norm_key(wo)] = {
            "wo": _norm_key(wo),
            "po": po,
            "company": company,
            "job_name": job_name,
            "cust_ship": due,
            "part_count": len(part_rows),
        }

    wip_wos = set(wip_summary_by_wo)
    sch_wos = set(sch_summary_by_wo)
    new_wos = wip_wos - sch_wos
    removed_wos = sch_wos - wip_wos

    if "MAIN" in wb_sch.sheetnames:
        del wb_sch["MAIN"]
    ws_out = wb_sch.create_sheet("MAIN", 0)
    _setup_sheet(ws_out)

    out_row = DATA
    wo_count = part_count = img_count = routing_count = 0
    routing_counts_by_wo = {}
    missing_routing_by_wo = {}
    ship_date_updates = []
    alt = False

    for wo, rows in wo_groups.items():
        wo_k = _norm_key(wo)
        wo_count += 1
        first_er, first_row = rows[0]
        first_is_summary = (first_row[0] is None and first_row[1] is None)

        # Header metadata
        job_name = (first_row[10] or first_row[9]) if first_is_summary else None
        if job_name: job_name = str(job_name).strip()
        company  = first_row[6]; po = first_row[7]; due = first_row[3]
        ship_date = ship_by_wo.get(wo_k, due)
        if isinstance(ship_date, datetime) and isinstance(due, datetime) and ship_date.date() != due.date():
            ship_date_updates.append({
                "wo": wo_k,
                "po": po,
                "company": company,
                "job_name": job_name,
                "google_ship": due,
                "excel_ship": ship_date,
            })

        # ── Navy header row ───────────────────────────────────────────────
        _write_header(ws_out, out_row, wo, po, company, job_name, ship_date)
        out_row += 1

        # ── Determine part rows ───────────────────────────────────────────
        part_rows = rows[1:] if first_is_summary else rows
        if not part_rows:
            part_rows = [(first_er, first_row)]   # single-row WO → also a part

        # ── Write part rows ───────────────────────────────────────────────
        for er, row_data in part_rows:
            part_count += 1
            alt = not alt

            # Old format col mapping (1-indexed in WIP, 0-indexed in row_data tuple)
            pno     = row_data[9]    # col J = PART NO (description for some WOs)
            desc    = row_data[10]   # col K = PART NAME (part# for some WOs)
            qty     = row_data[11]
            mat     = row_data[14]
            thk     = row_data[13]
            current = row_data[1]    # col B = CURRENT
            notes   = row_data[0]    # col A = ALL STEPS → reference notes
            ship    = row_data[3]    # col D = DUE DATE → delivery date

            # Pull routing + current/status/notes from the schedule file when matched
            routing = get_routing(wo, pno, by_part, by_wo)
            state = get_schedule_state(wo, pno, state_by_part, state_by_wo)
            current = state["current"] if state["current"] not in (None, "") else current
            status = state["status"]
            notes = state["notes"] if state["notes"] not in (None, "") else notes
            has_schedule_routing = bool(routing)

            # If no routing has delivery date, pre-fill from old due date
            if OUT_DELIVERY_COL not in routing and ship:
                routing[OUT_DELIVERY_COL] = ship

            if has_schedule_routing:
                routing_count += 1
                routing_counts_by_wo[wo_k] = routing_counts_by_wo.get(wo_k, 0) + 1
            else:
                missing_routing_by_wo[wo_k] = missing_routing_by_wo.get(wo_k, 0) + 1

            _write_part(ws_out, out_row, wo, po, company,
                        pno, desc, qty, mat, thk,
                        current, status, notes, routing, alt)

            # Screenshot: prefer WIP(1) source (most up-to-date)
            img_placed = False
            if er in wip_img_map:
                if _copy_image(wip_img_map[er], ws_out, 9, out_row):
                    img_count += 1; img_placed = True

            # Fallback: check Schedule(1) by key
            if not img_placed:
                wo_k  = _norm_key(wo)
                pno_k = _norm_key(pno)
                for key in [(wo_k, pno_k), (wo_k, None)]:
                    if key in sch_img_by_key:
                        if _copy_image(sch_img_by_key[key], ws_out, 9, out_row):
                            img_count += 1; break

            out_row += 1

    _apply_cf(ws_out, out_row - 1)
    _apply_dropdown(ws_out, out_row - 1)
    ws_out.freeze_panes = f"A{DATA}"
    ws_out.auto_filter.ref = f"A{HDR}:{LAST}{out_row - 1}"

    # The working file stays clean — the change report goes to its own file so
    # the reviewer has a diff to check instead of eyeballing the whole sheet.
    if "MERGE CHANGES" in wb_sch.sheetnames:
        del wb_sch["MERGE CHANGES"]

    print(f"\n  Result: {wo_count} WOs | {part_count} parts | "
          f"{routing_count} parts with routing | {img_count} screenshots")
    print(f"  Changes: {len(new_wos)} new WOs | {len(removed_wos)} removed WOs | {len(ship_date_updates)} ship date updates")

    write_review_report(wip_path, sch_path, wip_summary_by_wo, sch_summary_by_wo,
                        new_wos, removed_wos, ship_date_updates,
                        routing_counts_by_wo, missing_routing_by_wo)

    if DRY_RUN:
        print(f"\nDRY RUN — nothing written to {os.path.basename(OUT_FILE)}")
        print(f"Review {os.path.basename(REVIEW_FILE)}, then re-run without --dry-run.")
        return

    _backup_file(OUT_FILE)
    print(f"Saving → {OUT_FILE}")
    wb_sch.save(OUT_FILE)
    print("Done ✓\n")
    print(f"Next: review {os.path.basename(REVIEW_FILE)}, then run  python3 cmf_schedule_builder.py")


def write_review_report(wip_path, sch_path, wip_summary_by_wo, sch_summary_by_wo,
                        new_wos, removed_wos, ship_date_updates,
                        routing_counts_by_wo, missing_routing_by_wo):
    """Write the merge diff to a standalone workbook for the pre-push review step."""
    wb = openpyxl.Workbook()
    del wb[wb.sheetnames[0]]
    build_change_log_sheet(wb, wip_path, sch_path, wip_summary_by_wo, sch_summary_by_wo,
                           new_wos, removed_wos, ship_date_updates,
                           routing_counts_by_wo, missing_routing_by_wo)
    wb.save(REVIEW_FILE)
    print(f"  Review report → {os.path.basename(REVIEW_FILE)}")


if __name__ == "__main__":
    merge()
